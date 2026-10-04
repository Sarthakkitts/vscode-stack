import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

doc = docx.Document()
for section in doc.sections:
  section.top_margin = Inches(1)
  section.bottom_margin = Inches(1)
  section.left_margin = Inches(1)
  section.right_margin = Inches(1)


def style_table(table):
  table.alignment = WD_TABLE_ALIGNMENT.CENTER
  for i, row in enumerate(table.rows):
    for j, cell in enumerate(row.cells):
      cell.paragraphs[0].paragraph_format.space_before = Pt(4)
      cell.paragraphs[0].paragraph_format.space_after = Pt(4)
      if i == 0:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '1F4E78')
        tcPr.append(shd)
        for run in cell.paragraphs[0].runs:
          run.font.bold = True
          run.font.color.rgb = RGBColor(255, 255, 255)
      elif i % 2 == 1:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F2F4F7') 
        shd.set(qn('w:fill'),'F2F4F7')
        tcPr.append(shd)


# Save document function call after building content...
doc.save('SahaYatri_Technical_Specification.docx')
